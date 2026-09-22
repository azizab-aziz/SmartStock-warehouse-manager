"""
export_dispatch.py - reads the two scratch CSVs written by
db_export_do_items_csv / db_export_returns_csv and writes them as two
plain sheets in one .xlsx. No macros, no template - just openpyxl.

Usage:
    python export_dispatch.py <do_items_csv> <returns_csv> <xlsx_path> <label>
"""
import sys
import csv
import os

try:
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment
    from openpyxl.utils import get_column_letter
except ImportError:
    print("ERROR: openpyxl is not installed. Run: pip install openpyxl", file=sys.stderr)
    sys.exit(1)

HEADER_FILL = PatternFill(start_color="1E293B", end_color="1E293B", fill_type="solid")
HEADER_FONT = Font(color="FFFFFF", bold=True)


def read_csv(path):
    if not os.path.isfile(path):
        raise FileNotFoundError(f"CSV introuvable: {path}")
    with open(path, "r", encoding="utf-8-sig", newline="") as f:
        reader = csv.reader(f)
        rows = list(reader)
    if not rows:
        return [], []
    return rows[0], rows[1:]


def write_sheet(wb, sheet_name, header, rows, first=False):
    ws = wb.active if first else wb.create_sheet()
    ws.title = sheet_name[:31]

    for col_idx, col_name in enumerate(header, start=1):
        cell = ws.cell(row=1, column=col_idx, value=col_name)
        cell.font = HEADER_FONT
        cell.fill = HEADER_FILL
        cell.alignment = Alignment(horizontal="left")

    for row_idx, row in enumerate(rows, start=2):
        for col_idx, value in enumerate(row, start=1):
            ws.cell(row=row_idx, column=col_idx, value=value)

    widths = [len(str(h)) for h in header]
    for row in rows:
        for i, value in enumerate(row):
            if i < len(widths):
                widths[i] = max(widths[i], len(str(value)))
    for i, w in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(i)].width = min(max(w + 2, 10), 40)

    ws.freeze_panes = "A2"


def main():
    if len(sys.argv) != 5:
        print("ERROR: expected 4 arguments: do_items_csv returns_csv xlsx_path label", file=sys.stderr)
        sys.exit(1)

    do_csv_path, returns_csv_path, xlsx_path, label = sys.argv[1:5]

    try:
        do_header, do_rows = read_csv(do_csv_path)
        returns_header, returns_rows = read_csv(returns_csv_path)
    except FileNotFoundError as e:
        print(f"ERROR: {e}", file=sys.stderr)
        sys.exit(1)

    wb = Workbook()
    write_sheet(wb, f"{label} - Expeditions"[:31], do_header, do_rows, first=True)
    write_sheet(wb, f"{label} - Retours"[:31], returns_header, returns_rows, first=False)

    os.makedirs(os.path.dirname(xlsx_path) or ".", exist_ok=True)
    wb.save(xlsx_path)
    print(f"OK: {xlsx_path} genere ({len(do_rows)} lignes expeditions, {len(returns_rows)} lignes retours)")
    sys.exit(0)


if __name__ == "__main__":
    main()
